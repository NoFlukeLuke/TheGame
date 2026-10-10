#!/usr/bin/env python3
"""Turns the steer pass of tools/sim/trick-audit.js into the steering spreadsheet.

    python3 tools/sim/trick-steer-report.py STEER.json --out trick-steer.xlsx [--csv out.csv] [--json out.json]

The steer pass scores each Trick ALONE on the audit's boards. On every board the
player picks the hand that scores best with the Trick: as dealt, with one
neighbour swap, and with one discard (then the swap). It does that for the best
hand of each size (2 to 5 cards) and for the best hand of any size. The same
search with no Trick is the baseline, so every number here is "how much better
the best hand gets when you play for this Trick". The headline is the average
over the four hand sizes with a swap and a discard: a player who plays a spread
of hand lengths and steers each one.

What a Trick pays in (Score / Focus / Time / ...) and what its number rests on
come from trick-audit-report.py, so the two sheets agree.
"""
import argparse, gzip, importlib.util, json, math, os, statistics

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

_spec = importlib.util.spec_from_file_location('audit_report', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'trick-audit-report.py'))
AR = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(AR)
TIERS, TIER_WORD = AR.TIERS, AR.TIER_WORD
LEVELS = ['dealt', 'swap', 'discard']
LEVEL_WORD = {'dealt': 'as dealt', 'swap': 'one swap', 'discard': 'swap + discard'}


def lift(t, b, lv, k):
    """exp(mean over boards of log(best with the Trick) - log(best without)); boards where either has no hand of size k are skipped."""
    s = n = 0
    for v, w in zip(t['lv'][lv][k], b['lv'][lv][k]):
        if v is None or w is None: continue
        s += v - w; n += 1
    return (math.exp(s / n), n) if n else (None, 0)


def pays(t, lv, k):
    a = [v for v in t['fv'][lv][k] if v is not None]
    return sum(a) / len(a) if a else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('steer'); ap.add_argument('--out', required=True); ap.add_argument('--csv'); ap.add_argument('--json')
    a = ap.parse_args()
    d = json.load(gzip.open(a.steer, 'rt') if a.steer.endswith('.gz') else open(a.steer))
    cfg, pool = d['cfg'], d['pool']
    res = {r['id']: r for r in d['steer'] if r}
    base = res[None] if None in res else next(r for r in d['steer'] if r['id'] is None)
    sizes = [str(k) for k in range(2, cfg['maxHand'] + 1)]

    rows = []
    for t in pool:
        tid = t['id']; r = res.get(tid)
        row = dict(id=tid, name=t['name'], tier=t['tier'], desc=t['desc'], type=AR.ptype(tid), depends=AR.DEPENDS.get(tid, ''))
        if not r: rows.append(row); continue
        rn = r['rand']['n']
        row['rand_x'] = math.exp(r['rand']['log'] / rn) if rn else None
        row['rand_pays'] = r['rand']['fired'] / rn if rn else None
        for li, lv in enumerate(LEVELS):
            row[lv + '_x'] = lift(r, base, li, 'any')[0]
            row[lv + '_pays'] = pays(r, li, 'any')
        logs, pz = [], []
        for k in sizes:
            x, n = lift(r, base, 2, k)
            row['size' + k] = x
            row['size' + k + '_pays'] = pays(r, 2, k)
            if x: logs.append(math.log(x))
            if row['size' + k + '_pays'] is not None: pz.append(row['size' + k + '_pays'])
        row['all_x'] = math.exp(sum(logs) / len(logs)) if logs else None
        row['all_pays'] = sum(pz) / len(pz) if pz else None
        # How much of the steered value needs the swap and the discard.
        nb = len(r['lv'][0]['any'])
        row['swap_used'] = len(r['used']['swap']['any']) / nb if nb else None
        row['disc_used'] = len(r['used']['disc']['any']) / nb if nb else None
        row['names'] = r.get('names', {})
        row['err'] = r.get('err')
        rows.append(row)

    score = [r for r in rows if r['type'] == 'Score' and r.get('all_x')]
    # Rarity medians (the typical Trick of each tier; the mean is set by a few runaways).
    med = {}
    for t in TIERS:
        T = [r for r in score if r['tier'] == t]
        if not T: continue
        m = dict(n=len(T))
        for key in ['rand_x', 'dealt_x', 'swap_x', 'discard_x', 'all_x'] + ['size' + k for k in sizes]:
            v = [r[key] for r in T if r.get(key)]
            m[key] = statistics.median(v) if v else None
        for key in ['rand_pays', 'dealt_pays', 'swap_pays', 'discard_pays', 'all_pays']:
            v = [r[key] for r in T if r.get(key) is not None]
            m[key] = statistics.median(v) if v else None
        med[t] = m
    for r in score:
        m = med[r['tier']]['all_x']
        r['ratio'] = (r['all_x'] - 1) / (m - 1) if m and m > 1 else None
    # Suggested tier: rank the Score Tricks by their steered value and hand out
    # the tiers in today's numbers, weakest first (so the pool keeps its shape).
    order = sorted(score, key=lambda r: r['all_x'])
    counts = [sum(1 for r in score if r['tier'] == t) for t in TIERS]
    i = 0
    for t, c in zip(TIERS, counts):
        for r in order[i:i + c]: r['suggest'] = t
        i += c
    # Cut lines between the suggested tiers (the steered value at each boundary).
    cuts = []
    i = 0
    for t, c in zip(TIERS[:-1], counts[:-1]):
        i += c
        if 0 < i < len(order): cuts.append((t, order[i - 1]['all_x'], order[i]['all_x']))

    write_xlsx(a.out, rows, med, cuts, cfg, d, sizes, base)
    if a.csv:
        import csv
        cols = ['name', 'tier', 'type', 'suggest', 'ratio', 'all_x', 'all_pays', 'rand_x', 'dealt_x', 'swap_x', 'discard_x'] + ['size' + k for k in sizes] + ['desc']
        with open(a.csv, 'w', newline='') as f:
            w = csv.writer(f); w.writerow(cols)
            for r in sorted(rows, key=lambda r: (r['type'], -(r.get('all_x') or 0))):
                w.writerow([round(r[c], 3) if isinstance(r.get(c), float) else r.get(c) for c in cols])
    if a.json:
        json.dump(dict(rows=rows, med=med, cuts=cuts), open(a.json, 'w'))
    errs = [r['name'] for r in rows if r.get('err')]
    print('wrote', a.out, '| Score Tricks', len(score), '| errors in', errs[:8])


def write_xlsx(path, rows, med, cuts, cfg, d, sizes, base):
    wb = Workbook()
    hdr_fill = PatternFill('solid', fgColor='1F2A36'); hdr_font = Font(bold=True, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='top'); top = Alignment(vertical='top')
    move_fill = {1: 'FBE3B0', 2: 'F4B6B6', 3: 'F4B6B6', -1: 'DCE8F5', -2: 'C9D9EE', -3: 'C9D9EE'}

    def sheet(ws, headers, data, widths, notes=None):
        r0 = 1
        if notes:
            for n in notes:
                ws.cell(row=r0, column=1, value=n)
                r0 += 1
            r0 += 1
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=r0, column=j, value=h); c.fill = hdr_fill; c.font = hdr_font
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for i, rowv in enumerate(data, r0 + 1):
            for j, v in enumerate(rowv, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.alignment = wrap if isinstance(v, str) and len(v) > 24 else top
        for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(data)}"
        return r0

    def pct(v): return None if v is None else round(v * 100, 1)
    def rnd(v, k=2): return None if v is None else round(v, k)

    # ── 1. Steering ──
    ws = wb.active; ws.title = 'Steering'
    order = {'Score': 0, 'Focus': 1, 'Time': 2, 'Credits / stock': 3, 'Dead in Flow': 4, 'Not measured': 5}
    rs = sorted(rows, key=lambda r: (order.get(r['type'], 9), TIERS.index(r['tier']), -(r.get('all_x') or 0)))
    headers = ['Trick', 'Rarity', 'Pays in', 'Description',
               'Steered, all hand sizes (x)', 'Ratio to rarity median', 'Suggested rarity', 'Pays on, steered (% of boards)',
               'Random hand (x)', 'Random hand: pays on %',
               'Best hand, as dealt (x)', 'As dealt: pays on %', 'Best hand, one swap (x)', 'One swap: pays on %',
               'Best hand, swap + discard (x)', 'Swap + discard: pays on %'] + \
              [f'{k}-card hands, steered (x)' for k in sizes] + ['Swap used (% of boards)', 'Discard used (% of boards)', 'Rests on']
    data = []
    for r in rs:
        data.append([r['name'], TIER_WORD[r['tier']], r['type'], r['desc'],
                     rnd(r.get('all_x')), rnd(r.get('ratio')), TIER_WORD.get(r.get('suggest'), None), pct(r.get('all_pays')),
                     rnd(r.get('rand_x')), pct(r.get('rand_pays')),
                     rnd(r.get('dealt_x')), pct(r.get('dealt_pays')), rnd(r.get('swap_x')), pct(r.get('swap_pays')),
                     rnd(r.get('discard_x')), pct(r.get('discard_pays'))] +
                    [rnd(r.get('size' + k)) for k in sizes] + [pct(r.get('swap_used')), pct(r.get('disc_used')), r.get('depends', '')])
    notes = ['Each Trick ALONE. On 100 boards the player picks the hand that scores best with the Trick: as dealt, with one neighbour swap, and with one discard of up to 3 cards (then the swap).',
             '"x" = how much better that best hand scores than the best hand a player with no Trick finds the same way (x1.50 = 50% more).',
             '"Steered, all hand sizes" = the average over 2-, 3-, 4- and 5-card hands with a swap and a discard: a player who plays a spread of hand lengths and steers each one.',
             '"Ratio to rarity median": 1 = the typical Trick of its rarity. "Suggested rarity": rank every Score Trick by the steered number and hand out the rarities in today\'s numbers, weakest first.',
             'Only Score Tricks are ranked. Focus, Time and credit Tricks are listed for their score side only (most have none).']
    r0 = sheet(ws, headers, data, [22, 11, 10, 46, 12, 11, 11, 12, 10, 10, 10, 10, 10, 10, 11, 11] + [10] * len(sizes) + [10, 10, 28], notes)
    for i, r in enumerate(rs, r0 + 1):
        if r.get('suggest'):
            mv = TIERS.index(r['suggest']) - TIERS.index(r['tier'])
            if mv: ws.cell(row=i, column=7).fill = PatternFill('solid', fgColor=move_fill[max(-3, min(3, mv))])
    n = r0 + len(rs)
    ws.conditional_formatting.add(f'F{r0 + 1}:F{n}', ColorScaleRule(start_type='num', start_value=0, start_color='C9D9EE',
                                                                     mid_type='num', mid_value=1, mid_color='FFFFFF', end_type='num', end_value=4, end_color='F4B6B6'))

    # ── 2. By rarity ──
    ws = wb.create_sheet('By rarity')
    keys = [('rand_x', 'Random hand'), ('dealt_x', 'Best hand, as dealt'), ('swap_x', 'Best hand, one swap'), ('discard_x', 'Best hand, swap + discard')] + \
           [('size' + k, f'{k}-card hands, steered') for k in sizes] + [('all_x', 'Steered, all hand sizes')]
    data = []
    for key, label in keys:
        data.append([label] + [rnd(med[t].get(key)) if t in med else None for t in TIERS])
    for key, label in [('rand_pays', 'Pays on: random hand (%)'), ('dealt_pays', 'Pays on: as dealt (%)'), ('swap_pays', 'Pays on: one swap (%)'),
                       ('discard_pays', 'Pays on: swap + discard (%)'), ('all_pays', 'Pays on: steered, all sizes (%)')]:
        data.append([label] + [pct(med[t].get(key)) if t in med else None for t in TIERS])
    data.append(['Score Tricks'] + [med[t]['n'] if t in med else None for t in TIERS])
    sheet(ws, ['Median Score Trick', 'Common', 'Rare', 'Epic', 'Legendary'], data, [34, 12, 12, 12, 12],
          ['The typical (median) Score Trick of each rarity, each one alone. The median, not the average: a few runaways set the average.'])

    # ── 3. Moves ──
    ws = wb.create_sheet('Suggested moves')
    score = [r for r in rows if r.get('suggest')]
    data = []
    for t in TIERS:
        data.append([TIER_WORD[t]] + [sum(1 for r in score if r['tier'] == t and r['suggest'] == s) for s in TIERS])
    r0 = sheet(ws, ['Today \\ Suggested', 'Common', 'Rare', 'Epic', 'Legendary'], data, [26, 12, 12, 12, 12],
               ['Score Tricks by today\'s rarity (rows) and the rarity their steered value ranks into (columns). The diagonal stays put.',
                'Cut lines (steered, all sizes): ' + '; '.join(f'{TIER_WORD[t]} up to x{lo:.2f}, next from x{hi:.2f}' for t, lo, hi in cuts)])
    moves = sorted([r for r in score if r['suggest'] != r['tier']],
                   key=lambda r: (-abs(TIERS.index(r['suggest']) - TIERS.index(r['tier'])), -r['all_x']))
    ws2 = wb.create_sheet('Moves list')
    sheet(ws2, ['Trick', 'Today', 'Suggested', 'Steps', 'Steered, all sizes (x)', 'Random hand (x)', 'Description'],
          [[r['name'], TIER_WORD[r['tier']], TIER_WORD[r['suggest']], TIERS.index(r['suggest']) - TIERS.index(r['tier']),
            rnd(r['all_x']), rnd(r.get('rand_x')), r['desc']] for r in moves], [22, 11, 11, 7, 12, 12, 60],
          ['Every Score Trick whose steered value ranks into another rarity. Positive steps = stronger than its rarity (move up or tune down).'])

    # ── 4. Settings ──
    ws = wb.create_sheet('Settings')
    explain = {'scoring': 'Scoring model (classic / mult_ladder / hand_size)', 'steerGrids': 'Boards searched per Trick',
               'steerSamples': 'Random draws behind each discard', 'steerCands': 'Discard groups tried per board',
               'steerDiscardMax': 'Most cards in one discard', 'level': 'Flow level', 'quarter': 'Quarter (QRL)', 'maxHand': 'Biggest hand',
               'rows': 'Board rows', 'cols': 'Board columns', 'seed': 'Random seed', 'handSeconds': 'Average seconds between hands',
               'holdHands': 'Hands a scaling Trick has been held', 'sleightsOwned': 'Sleights owned', 'coinsMax': 'Credits held: 0 to this'}
    data = [[k, cfg[k] if not isinstance(cfg[k], (list, dict)) else 'from the Focus run', explain.get(k, '')] for k in cfg if k != 'focusDist']
    data.append(['build', d.get('build', ''), 'Game build the run used'])
    data.append(['started', d.get('started', ''), ''])
    sheet(ws, ['Setting', 'Value', 'What it is'], data, [22, 22, 60],
          ['The discard tried on each board is the one (of the least disruptive groups) whose random draws give the best hands across sizes;',
           'a size takes it only when its average beats not discarding. A used swap or discard counts for Tricks that read them; its clock cost does not.',
           'A Trick that pays on small or fast hands can look weaker here than it plays: the player picks the highest score, not the quickest hand.'])
    wb.save(path)


if __name__ == '__main__':
    main()
