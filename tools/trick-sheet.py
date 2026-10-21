#!/usr/bin/env python3
"""The Trick balance sheet (r620).

    node tools/trick-sheet.js common
    python3 tools/trick-sheet.py tools/sim/out/tricks-common.json --out TRICKS_COMMON.xlsx [--audit audit.csv] [--csv out.csv]

One row per NUMBER a Trick has in BAL (a Trick with two numbers takes two rows;
one with none takes one row that says so). Yellow columns are the owner's to
fill: the new value, a verdict, notes. The measured columns come from
tools/sim/trick-audit.js through trick-audit-report.py --csv (matched by name).
"""
import argparse, csv, json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument('json'); ap.add_argument('--out', required=True); ap.add_argument('--audit'); ap.add_argument('--csv')
a = ap.parse_args()
d = json.load(open(a.json))
audit = {}
if a.audit:
    for r in csv.DictReader(open(a.audit)):
        audit[r['name']] = r

HEAD = ['#', 'Trick', 'id', 'Rarity', 'Category', 'What it does (live text)', 'Number', 'Now', 'NEW VALUE', 'x factor',
        'Share of score', 'x on a hand', 'Fires %', 'Audit note', 'Verdict', 'Notes', 'Template', 'Tags']
WIDTH = [4, 20, 20, 10, 16, 56, 18, 8, 11, 9, 10, 10, 8, 30, 10, 40, 50, 22]
EDIT = {'NEW VALUE', 'Verdict', 'Notes'}

wb = Workbook(); ws = wb.active; ws.title = d['tier'].capitalize() + ' Tricks'
yellow = PatternFill('solid', fgColor='FFF4C2'); grey = PatternFill('solid', fgColor='EFEFEF'); head = PatternFill('solid', fgColor='2B2440')
thin = Side(style='thin', color='CCCCCC')
ws.append(HEAD)
for i, h in enumerate(HEAD, 1):
    c = ws.cell(row=1, column=i); c.font = Font(bold=True, color='FFFFFF'); c.fill = head
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = WIDTH[i - 1]
ws.row_dimensions[1].height = 30

def num(v):
    try: return round(float(v), 2)
    except (TypeError, ValueError): return None

row = 2; n = 0
for t in d['tricks']:
    n += 1
    au = audit.get(t['name'], {})
    share = num(au.get('ratio_any')); xh = num(au.get('x_any')); fires = num(au.get('fire_any'))
    note = au.get('flag', '') if au.get('flag', '-') != '-' else ''
    if au and au.get('type') and au['type'] != 'Score': note = (note + ' ' if note else '') + f"pays in {au['type']}"
    nums = t['nums'] or [{'key': '(in code, no BAL number)', 'value': None}]
    for j, kv in enumerate(nums):
        first = j == 0
        vals = [n if first else None, t['name'] if first else '', t['id'] if first else '', t['tier'] if first else '', t['category'] if first else '',
                t['desc'] if first else '', kv['key'], kv['value'], None, None,
                share if first else None, xh if first else None, (round(fires * 100) if fires is not None else None) if first else None,
                note if first else '', '', '', t['template'] if first else '', ', '.join(t['tags']) if first else '']
        ws.append(vals)
        if kv['value'] is not None:
            ws.cell(row=row, column=10).value = f'=IF(I{row}="","",I{row}/H{row})'
        for col in range(1, len(HEAD) + 1):
            c = ws.cell(row=row, column=col)
            c.alignment = Alignment(wrap_text=col in (6, 14, 16, 17), vertical='top')
            c.border = Border(bottom=thin)
            if HEAD[col - 1] in EDIT: c.fill = yellow
            elif not first: c.fill = grey
        if fires is not None and first: ws.cell(row=row, column=13).number_format = '0'
        for col in (10, 11, 12): ws.cell(row=row, column=col).number_format = '0.00'
        row += 1

dv = DataValidation(type='list', formula1='"KEEP,UP,EDIT,CUT"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'O2:O{row - 1}')
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(HEAD))}{row - 1}'

# ── the key ──
k = wb.create_sheet('Read me')
lines = [
    ('Trick balance sheet', True), (f"Build: {d['build']}", False), (f"Tier: {d['tier']} · {len(d['tricks'])} Tricks", False), ('', False),
    ('Yellow columns are yours. Everything else is read from the game as it ships.', False), ('', False),
    ('Number / Now', True), ('The Trick\'s tunable in BAL (js/data/balance.js) and its shipped value. A Trick with two numbers has two rows. "(in code)" means the value is typed into the scoring code, not BAL: changing it is a code edit.', False),
    ('NEW VALUE / x factor', True), ('Type the new number. x factor works itself out (new / now). To scale the whole tier the same way, type the factor into a spare cell and set NEW VALUE to =H2*that.', False),
    ('Share of score', True), ('From the Trick audit sim (tools/sim/trick-audit.js, card ladder scoring, Flow level 12): how much of the score a random 5-Trick loadout owes to this Trick, as a multiple of the average Trick. 1.0 = an average Trick; 2.0 = twice the average; 0.3 = barely felt. Blank = the audit does not measure it (it pays in Focus, time or credits, or needs something the sim has no model of).', False),
    ('x on a hand', True), ('How much bigger a hand is with the Trick than without it, averaged over every hand the bot played (1.00 = no change). Fires % is the share of hands it paid on at all.', False),
    ('Verdict', True), ('KEEP / UP / EDIT / CUT. UP means "the number goes up, nothing else changes".', False),
    ('Trick tour', True), ('index.html?tricks opens the common Tricks one by one on a live board; its number boxes write the same BAL value, so a new number can be felt before it goes in the sheet. COPY ALL there prints old -> new per Trick.', False),
    ('Applying the sheet', True), ('Hand the sheet back; the numbers go into js/data/balance.js (BAL) by id and key, the descriptions regenerate from DESC_TEMPLATES.', False),
]
for i, (text, bold) in enumerate(lines, 1):
    c = k.cell(row=i, column=1, value=text); c.font = Font(bold=bold, size=12 if bold else 11); c.alignment = Alignment(wrap_text=True, vertical='top')
k.column_dimensions['A'].width = 110

wb.save(a.out)
if a.csv:
    with open(a.csv, 'w', newline='') as f:
        w = csv.writer(f); w.writerow(HEAD)
        for r in ws.iter_rows(min_row=2, values_only=True): w.writerow(['' if v is None or (isinstance(v, str) and v.startswith('=')) else v for v in r])
print(f'wrote {a.out}: {row - 2} rows, audit matched {sum(1 for t in d["tricks"] if t["name"] in audit)} of {len(d["tricks"])}')
